import sys
import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_review2_presentation(output_path):
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    PRIMARY_DARK = RGBColor(15, 23, 42)      # Deep Navy #0F172A
    ACCENT_BLUE  = RGBColor(30, 58, 138)     # Royal Blue #1E3A8A
    ACCENT_TEAL  = RGBColor(13, 148, 136)    # Teal #0D9488
    CARD_BG      = RGBColor(241, 245, 249)   # Light Slate #F1F5F9
    CARD_BORDER  = RGBColor(203, 213, 225)   # Border #CBD5E1
    TEXT_DARK    = RGBColor(15, 23, 42)      # Main Text #0F172A
    TEXT_MUTED   = RGBColor(71, 85, 105)     # Subtitle/Muted #475569
    TEXT_LIGHT   = RGBColor(255, 255, 255)   # White
    SUCCESS_GRN  = RGBColor(22, 101, 52)     # Dark Green #166534
    ACCENT_CYAN  = RGBColor(2, 132, 199)     # Sky Blue #0284C7

    def add_header(slide, title_text, category="PROJECT REVIEW - 2"):
        # Top Header Bar
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = PRIMARY_DARK
        header_box.line.color.rgb = PRIMARY_DARK

        # Category Tag
        cat_p = header_box.text_frame.paragraphs[0]
        cat_p.text = category.upper()
        cat_p.font.name = "Calibri"
        cat_p.font.size = Pt(11)
        cat_p.font.bold = True
        cat_p.font.color.rgb = ACCENT_TEAL
        cat_p.space_after = Pt(2)

        # Title
        title_p = header_box.text_frame.add_paragraph()
        title_p.text = title_text
        title_p.font.name = "Calibri"
        title_p.font.size = Pt(22)
        title_p.font.bold = True
        title_p.font.color.rgb = TEXT_LIGHT

        # Bottom Accent Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06))
        line.fill.solid()
        line.fill.fore_color.rgb = ACCENT_TEAL
        line.line.color.rgb = ACCENT_TEAL

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.733), Inches(0.35))
        tf = footer_box.text_frame
        p = tf.paragraphs[0]
        p.text = "Number System Toolkit | Department of CSE (Data Science) | Aditya University"
        p.font.name = "Calibri"
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, badge=""):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

        # Card Title
        p0 = tf.paragraphs[0]
        p0.text = f"{title}" if not badge else f"[{badge}]  {title}"
        p0.font.name = "Calibri"
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_BLUE
        p0.space_after = Pt(8)

        for itm in items:
            p = tf.add_paragraph()
            p.text = f"•  {itm}"
            p.font.name = "Calibri"
            p.font.size = Pt(11.5)
            p.font.color.rgb = TEXT_DARK
            p.space_after = Pt(4)

    # ----------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PRIMARY_DARK
    bg1.line.color.rgb = PRIMARY_DARK

    # Top accent pill
    pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(3.2), Inches(0.45))
    pill.fill.solid()
    pill.fill.fore_color.rgb = ACCENT_TEAL
    pill.line.color.rgb = ACCENT_TEAL
    ptf = pill.text_frame
    ptf.text = "PROJECT REVIEW - II"
    pp = ptf.paragraphs[0]
    pp.font.bold = True
    pp.font.size = Pt(13)
    pp.font.color.rgb = TEXT_LIGHT
    pp.alignment = PP_ALIGN.CENTER

    # Main Title Box
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(11.333), Inches(2.0))
    ttf = t_box.text_frame
    tp = ttf.paragraphs[0]
    tp.text = "NUMBER SYSTEM TOOLKIT"
    tp.font.name = "Arial Black"
    tp.font.size = Pt(40)
    tp.font.bold = True
    tp.font.color.rgb = TEXT_LIGHT

    sub_p = ttf.add_paragraph()
    sub_p.text = "Comprehensive Digital Logic, Base Conversions & Advanced Mathematical Suite"
    sub_p.font.name = "Calibri"
    sub_p.font.size = Pt(18)
    sub_p.font.color.rgb = RGBColor(148, 163, 184)
    sub_p.space_before = Pt(8)

    # Left Team Box
    tm_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(3.6), Inches(5.8), Inches(3.2))
    tm_card.fill.solid()
    tm_card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    tm_card.line.color.rgb = RGBColor(51, 65, 85)
    tm_tf = tm_card.text_frame
    tm_tf.margin_left = Inches(0.3)
    tm_tf.margin_top = Inches(0.25)
    tp1 = tm_tf.paragraphs[0]
    tp1.text = "TEAM MEMBERS"
    tp1.font.name = "Calibri"
    tp1.font.bold = True
    tp1.font.size = Pt(14)
    tp1.font.color.rgb = ACCENT_TEAL
    tp1.space_after = Pt(10)

    members = [
        ("25B11DS349", "M. LAKSHMI SRINIVAS"),
        ("25B11DS127", "D. VARAPRASAD"),
        ("25B11DS370", "N. AKHILA DEVI"),
        ("25B11DS628", "Y. NITHYA SREE"),
        ("25B11DS320", "M. VENKATA PRANATHI"),
    ]
    for roll, name in members:
        p = tm_tf.add_paragraph()
        p.text = f"{roll}   •   {name}"
        p.font.name = "Calibri"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(5)

    # Right Dept Box
    dp_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(3.6), Inches(5.2), Inches(3.2))
    dp_card.fill.solid()
    dp_card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    dp_card.line.color.rgb = RGBColor(51, 65, 85)
    dp_tf = dp_card.text_frame
    dp_tf.margin_left = Inches(0.3)
    dp_tf.margin_top = Inches(0.25)
    dp1 = dp_tf.paragraphs[0]
    dp1.text = "INSTITUTION & DEPARTMENT"
    dp1.font.name = "Calibri"
    dp1.font.bold = True
    dp1.font.size = Pt(14)
    dp1.font.color.rgb = ACCENT_TEAL
    dp1.space_after = Pt(12)

    lines = [
        "Department of Computer Science & Engineering",
        "(Data Science)",
        "",
        "ADITYA UNIVERSITY",
        "Academic Year: 2025 - 2026",
        "Stage: Review - II Presentation & Prototype Demo"
    ]
    for l in lines:
        p = dp_tf.add_paragraph()
        p.text = l
        p.font.name = "Calibri"
        p.font.size = Pt(12 if l != "ADITYA UNIVERSITY" else 14)
        p.font.bold = (l == "ADITYA UNIVERSITY" or l == "Department of Computer Science & Engineering")
        p.font.color.rgb = TEXT_LIGHT if l != "(Data Science)" else ACCENT_CYAN
        p.space_after = Pt(2)

    # ----------------------------------------------------
    # SLIDE 2: REVIEW 1 COMMITMENTS VS REVIEW 2 STATUS
    # ----------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Review 1 Commitments vs. Review 2 Deliverables")

    add_card(s2, 0.8, 1.4, 2.7, 5.3, "1. Finalize Requirements", [
        "Status: COMPLETED (100%)",
        "Identified standard number bases: Binary, Octal, Decimal, Hexadecimal, and Custom Base (2-36).",
        "Added ASCII text encoding & decoding.",
        "Included 1's & 2's complement and full bitwise arithmetic."
    ], "PHASE 1")

    add_card(s2, 3.8, 1.4, 2.7, 5.3, "2. Interface Design", [
        "Status: COMPLETED (100%)",
        "Modern responsive Glassmorphic web UI designed.",
        "Categorized tabs for Conversions, Step-by-Step Solver, Bitwise Board, Trig Lab, and Log Suite.",
        "Interactive canvas for visual unit circle representation."
    ], "PHASE 2")

    add_card(s2, 6.8, 1.4, 2.7, 5.3, "3. Toolkit Development", [
        "Status: COMPLETED (100%)",
        "C++ Core calculation logic built & modularized.",
        "Full client-side JavaScript engine for real-time instant computations.",
        "Step-by-step mathematical division & expansion derivation generator."
    ], "PHASE 3")

    add_card(s2, 9.8, 1.4, 2.7, 5.3, "4. Prototype Testing", [
        "Status: COMPLETED (100%)",
        "Verified all conversion algorithms against manual calculations.",
        "Validated edge cases: zero, negative inputs, large numbers, fraction bases.",
        "Integrated interactive challenge quiz for validation."
    ], "PHASE 4")

    # ----------------------------------------------------
    # SLIDE 3: SYSTEM ARCHITECTURE & DESIGN
    # ----------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "System Architecture & Modular Design")

    add_card(s3, 0.8, 1.4, 3.6, 5.3, "1. Presentation / UI Layer", [
        "Interactive Web Interface (HTML5, Modern CSS Glassmorphism).",
        "Dynamic Unit Circle Canvas (HTML5 2D Canvas API).",
        "Interactive 8-Bit manipulation switchboard.",
        "Cheat Sheets, Lookup Tables & Practice Quiz UI.",
        "Light & Dark Mode adaptive themes."
    ], "FRONTEND")

    add_card(s3, 4.8, 1.4, 3.6, 5.3, "2. Core Calculation Engine", [
        "Base Converter: Repeated division & positional power expansion algorithms.",
        "Binary Arithmetic: Full adder logic, signed subtraction, multiplication & division.",
        "Bitwise Processor: AND, OR, XOR, NOT, Left/Right bit shifts, 2's complement.",
        "Trigonometric & Logarithmic high-precision math evaluators."
    ], "ENGINE")

    add_card(s3, 8.8, 1.4, 3.6, 5.3, "3. Educational & Step Solvers", [
        "Step-by-Step Derivation Generator: shows full intermediate steps with quotient & remainder.",
        "Positional Value Breakdown: generates weighted summation equations.",
        "Trig Function Visualizer: real-time sine/cosine projection on Cartesian grid.",
        "Instant Quiz Engine: random challenge generator with instant score tracking."
    ], "LEARNING")

    # ----------------------------------------------------
    # SLIDE 4: CORE MODULE 1 - MULTI-BASE CONVERTER & SOLVER
    # ----------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Module 1: Multi-Base Converter & Step-by-Step Solver")

    add_card(s4, 0.8, 1.4, 5.6, 5.3, "Supported Base Systems & Features", [
        "Binary (Base-2), Octal (Base-8), Decimal (Base-10), Hexadecimal (Base-16).",
        "Arbitrary Custom Base Conversion (from Base-2 up to Base-36).",
        "ASCII Text to Binary/Hex/Octal encoding and decoding.",
        "Dual input mode: real-time instantaneous updates across all bases.",
        "Input validation: prevents invalid characters (e.g. '8' in Octal, '2' in Binary)."
    ], "CAPABILITIES")

    add_card(s4, 6.8, 1.4, 5.6, 5.3, "Step-by-Step Mathematical Derivations", [
        "Repeated Division Method (Decimal to Base-N):",
        "  - Displays step-by-step division, quotient, and remainder table.",
        "  - Visualizes reading remainders from bottom to top (MSB to LSB).",
        "Positional Expansion Method (Base-N to Decimal):",
        "  - Breaks number into individual digit weights (e.g., d * base^p).",
        "  - Generates full algebraic summation string and final evaluated sum."
    ], "ALGORITHMS")

    # ----------------------------------------------------
    # SLIDE 5: CORE MODULE 2 - BINARY ARITHMETIC & BITWISE LAB
    # ----------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Module 2: Binary Arithmetic & Bitwise Logic Lab")

    add_card(s5, 0.8, 1.4, 5.6, 5.3, "Binary Arithmetic Operations", [
        "Binary Addition: full bit-by-bit addition with carry propagation.",
        "Binary Subtraction: standard borrowing logic and 2's complement subtraction.",
        "Binary Multiplication: shift-and-add binary partial product algorithm.",
        "Binary Division: restoring division algorithm with quotient and remainder.",
        "Provides simultaneous Decimal and Binary output representation."
    ], "ARITHMETIC")

    add_card(s5, 6.8, 1.4, 5.6, 5.3, "Bitwise Operations & Complement Suite", [
        "Logical Operators: Bitwise AND (&), OR (|), XOR (^), NOT (~).",
        "Bit Shifts: Left Shift (<<) and Right Shift (>>).",
        "1's Complement: Inversion of all bit positions.",
        "2's Complement: 1's complement + 1 with sign-bit preservation.",
        "Interactive 8-Bit Board: clickable individual bit toggles that update decimal/hex/char values dynamically in real-time."
    ], "BITWISE LAB")

    # ----------------------------------------------------
    # SLIDE 6: EXTENDED MODULE 3 - TRIGONOMETRY & LOGARITHM SUITE
    # ----------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Module 3: Advanced Mathematical Suites (Trig & Log)")

    add_card(s6, 0.8, 1.4, 5.6, 5.3, "Trigonometry Lab & Visual Unit Circle", [
        "Computes all 6 fundamental trigonometric functions:",
        "  - sin(θ), cos(θ), tan(θ), csc(θ), sec(θ), cot(θ).",
        "Supports both Degree (°) and Radian (rad) angle modes.",
        "Interactive Canvas Unit Circle:",
        "  - Dynamic ray rendering with live angle slider.",
        "  - Visual color-coded projections of Sine (vertical) and Cosine (horizontal) lengths.",
        "Handles undefined values gracefully (e.g. tan(90°), cot(0°))."
    ], "TRIGONOMETRY")

    add_card(s6, 6.8, 1.4, 5.6, 5.3, "Logarithm Suite & Change of Base", [
        "Natural Logarithm: ln(x) with base e (Euler's number).",
        "Common Logarithm: log10(x) for scientific & decibel calculations.",
        "Binary Logarithm: log2(x) essential for computer science & algorithmic complexity.",
        "Arbitrary Base Logarithm: log_b(x) using change of base formula:",
        "  - Formula: log_b(x) = ln(x) / ln(b).",
        "Exponentiation & power operations with high numeric precision."
    ], "LOGARITHMS")

    # ----------------------------------------------------
    # SLIDE 7: EDUCATIONAL LEARNING LAB & CHEAT SHEETS
    # ----------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Educational Hub: Cheat Sheets & Challenge Quizzes")

    add_card(s7, 0.8, 1.4, 5.6, 5.3, "Built-in Quick Reference Tables", [
        "Powers of 2 Lookup Table (2^0 to 2^16) with memory representations (1 KB, 2 KB, etc.).",
        "Full Number System Equivalence Matrix (0 to 15 in Decimal, Binary, Octal, and Hexadecimal).",
        "Common ASCII Character Codes & Hex/Binary mappings.",
        "Trigonometric values for standard angles (0°, 30°, 45°, 60°, 90°, 180°, 270°, 360°).",
        "Logarithmic identities and laws summary card."
    ], "REFERENCE")

    add_card(s7, 6.8, 1.4, 5.6, 5.3, "Interactive Challenge & Practice Quiz", [
        "Dynamic question generator across 3 difficulty tiers (Easy, Medium, Hard).",
        "Covers base conversion, binary addition, and bitwise logic puzzles.",
        "Instant automatic answer verification with detailed explanation modal.",
        "Live score tracker, streak counter, and retry mechanism for active learning.",
        "Helps university students practice for Digital Logic Design (DLD) exams."
    ], "PRACTICE QUIZ")

    # ----------------------------------------------------
    # SLIDE 8: TESTING, VERIFICATION & TEST CASES
    # ----------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Testing & Experimental Validation")

    # Table for Test Cases
    rows, cols = 7, 5
    t_shape = s8.shapes.add_table(rows, cols, Inches(0.8), Inches(1.4), Inches(11.733), Inches(4.5))
    table = t_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(2.2)
    table.columns[2].width = Inches(2.8)
    table.columns[3].width = Inches(2.8)
    table.columns[4].width = Inches(1.5)

    headers = ["Test Module", "Test Input", "Expected Output", "Actual Output", "Result"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY_DARK
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = "Calibri"
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT

    test_data = [
        ("Base Conversion", "Decimal: 255 to Hex / Bin", "Hex: FF | Bin: 11111111", "Hex: FF | Bin: 11111111", "PASSED [✓]"),
        ("Custom Base", "Decimal: 100 to Base-7", "202 (Base 7)", "202 (Base 7)", "PASSED [✓]"),
        ("Binary Addition", "10110 + 01101", "100011 (Decimal 35)", "100011 (Decimal 35)", "PASSED [✓]"),
        ("Bitwise XOR", "12 ^ 10 (1100 ^ 1010)", "0110 (Decimal 6)", "0110 (Decimal 6)", "PASSED [✓]"),
        ("Trig Function", "sin(30°) in Degree mode", "0.5000", "0.5000", "PASSED [✓]"),
        ("Logarithm", "log2(256) and log10(1000)", "8.0 and 3.0", "8.0 and 3.0", "PASSED [✓]"),
    ]

    for i, row in enumerate(test_data):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(248, 250, 252) if i % 2 == 0 else RGBColor(241, 245, 249)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = "Calibri"
            p.font.size = Pt(11)
            p.font.bold = (j == 0 or j == 4)
            if j == 4:
                p.font.color.rgb = SUCCESS_GRN
            else:
                p.font.color.rgb = TEXT_DARK

    # ----------------------------------------------------
    # SLIDE 9: TECHNOLOGY STACK & IMPLEMENTATION ENVIRONMENT
    # ----------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Technical Stack & Implementation Tools")

    add_card(s9, 0.8, 1.4, 3.6, 5.3, "C++ Core Engine", [
        "Language: C++ (C++17 standard).",
        "Algorithms: High efficiency integer arithmetic, recursive division, Bitset manipulation.",
        "Modular OOP structure: NumberSystem class with static conversion routines.",
        "Zero external dependencies for maximum cross-platform portability."
    ], "BACKEND / LOGIC")

    add_card(s9, 4.8, 1.4, 3.6, 5.3, "Modern Web UI Stack", [
        "HTML5: Semantic markup structure.",
        "CSS3: Glassmorphism UI, CSS Grid / Flexbox layout, dark/light theme variables.",
        "Vanilla JavaScript (ES6+): Event-driven reactive computation, zero heavy library overhead.",
        "HTML5 Canvas API: High performance math rendering for the Trigonometric unit circle."
    ], "FRONTEND UI")

    add_card(s9, 8.8, 1.4, 3.6, 5.3, "Development & QA Tools", [
        "IDE / Editor: Visual Studio Code, Antigravity Workspace.",
        "Compiler: GCC / G++ MinGW (for C++ engine).",
        "Testing: Automated browser test assertions and mathematical boundary validation.",
        "Version Control: Git version control for collaborative development."
    ], "TOOLING")

    # ----------------------------------------------------
    # SLIDE 10: INDIVIDUAL TEAM WORK DISTRIBUTION
    # ----------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Individual Team Work Distribution & Contributions")

    team_roles = [
        ("M. LAKSHMI SRINIVAS", "25B11DS349", "Team Lead & Core Algorithm Developer", [
            "Architected the modular C++ NumberSystem class structure.",
            "Designed and implemented multi-base conversion algorithms (Base 2-36).",
            "Coordinated overall module integration and Review 2 documentation."
        ]),
        ("D. VARAPRASAD", "25B11DS127", "Binary Arithmetic & Bitwise Logic Engineer", [
            "Implemented Binary Addition, Subtraction, Multiplication & Division logic.",
            "Developed bitwise manipulation logic (AND, OR, XOR, NOT, Shifts).",
            "Created 1's and 2's complement computation engine."
        ]),
        ("N. AKHILA DEVI", "25B11DS370", "Advanced Math Modules (Trig & Log)", [
            "Formulated Trigonometric computation logic for all 6 functions.",
            "Implemented logarithmic suites (ln, log10, log2, custom base change).",
            "Assisted in unit circle angle-to-coordinate mapping algorithms."
        ]),
        ("Y. NITHYA SREE", "25B11DS628", "UI/UX & Interactive Learning Hub", [
            "Designed modern Glassmorphic web interface layout and CSS styling.",
            "Built quick-reference Cheat Sheet tables and lookup matrices.",
            "Developed the interactive 8-Bit clickable switchboard UI."
        ]),
        ("M. VENKATA PRANATHI", "25B11DS320", "Frontend Canvas, Step-Solver & QA Testing", [
            "Engineered the dynamic HTML5 Canvas Unit Circle graphical visualizer.",
            "Implemented step-by-step mathematical division & power expansion generator.",
            "Designed comprehensive test suites and verified edge case accuracy."
        ])
    ]

    # Two rows for 5 members: 3 in row 1, 2 in row 2
    for idx, (name, roll, role, tasks) in enumerate(team_roles):
        if idx < 3:
            left = 0.8 + idx * 4.0
            top = 1.4
            w = 3.8
            h = 2.6
        else:
            left = 2.8 + (idx - 3) * 4.0
            top = 4.2
            w = 3.8
            h = 2.6

        card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_top = Inches(0.12)
        tf.margin_right = Inches(0.15)

        p = tf.paragraphs[0]
        p.text = f"{name} ({roll})"
        p.font.name = "Calibri"
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = PRIMARY_DARK

        p_r = tf.add_paragraph()
        p_r.text = role
        p_r.font.name = "Calibri"
        p_r.font.bold = True
        p_r.font.size = Pt(10.5)
        p_r.font.color.rgb = ACCENT_TEAL
        p_r.space_after = Pt(3)

        for t in tasks:
            pt = tf.add_paragraph()
            pt.text = f"• {t}"
            pt.font.name = "Calibri"
            pt.font.size = Pt(9.5)
            pt.font.color.rgb = TEXT_MUTED

    # ----------------------------------------------------
    # SLIDE 11: CONCLUSION & ROADMAP FOR REVIEW 3
    # ----------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Conclusion & Roadmap for Review 3 (Final Phase)")

    add_card(s11, 0.8, 1.4, 5.6, 5.3, "Summary of Review 2 Milestones", [
        "Successfully met 100% of the objectives committed during Review 1.",
        "Built a complete dual-architecture product: high-performance C++ Core and rich interactive Web Application.",
        "Features comprehensive multi-base conversion with full step-by-step mathematical reasoning.",
        "Added advanced interactive mathematical suites (Trig Unit Circle and Log Suite).",
        "Passed all verification test cases with zero calculation discrepancies."
    ], "ACHIEVEMENTS")

    add_card(s11, 6.8, 1.4, 5.6, 5.3, "Roadmap for Final Review (Review 3)", [
        "1. Performance Optimization & Big-Number Support:",
        "    - Support arbitrarily large bit-widths (32-bit, 64-bit, 128-bit).",
        "2. IEEE-754 Floating Point Representation:",
        "    - Single & double precision 32/64-bit floating point breakdown.",
        "3. Deployment & Live Hosting:",
        "    - Host the web application on GitHub Pages / Vercel for public access.",
        "4. Final Project Documentation & Viva Preparation:",
        "    - Prepare complete IEEE standard project report and user manual."
    ], "NEXT TARGETS")

    # ----------------------------------------------------
    # SLIDE 12: THANK YOU SLIDE
    # ----------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    bg12 = s12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = PRIMARY_DARK
    bg12.line.color.rgb = PRIMARY_DARK

    th_box = s12.shapes.add_textbox(Inches(2.0), Inches(2.2), Inches(9.333), Inches(3.0))
    th_tf = th_box.text_frame
    th_p = th_tf.paragraphs[0]
    th_p.text = "THANK YOU!"
    th_p.font.name = "Arial Black"
    th_p.font.size = Pt(48)
    th_p.font.bold = True
    th_p.font.color.rgb = TEXT_LIGHT
    th_p.alignment = PP_ALIGN.CENTER

    th_sub = th_tf.add_paragraph()
    th_sub.text = "Questions & Feedback Welcome"
    th_sub.font.name = "Calibri"
    th_sub.font.size = Pt(22)
    th_sub.font.color.rgb = ACCENT_TEAL
    th_sub.alignment = PP_ALIGN.CENTER
    th_sub.space_before = Pt(12)

    th_proj = th_tf.add_paragraph()
    th_proj.text = "NUMBER SYSTEM TOOLKIT  |  REVIEW - II"
    th_proj.font.name = "Calibri"
    th_proj.font.size = Pt(14)
    th_proj.font.color.rgb = RGBColor(148, 163, 184)
    th_proj.alignment = PP_ALIGN.CENTER
    th_proj.space_before = Pt(24)

    # Save
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\pranathi\.gemini\antigravity\scratch\number-system-toolkit"
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "Review_2_Number_System_Toolkit.pptx")
    create_review2_presentation(out_file)
